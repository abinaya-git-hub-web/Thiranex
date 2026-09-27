"""
=============================================================================
Thiranex Solutions — Multi-Format Report Generator Module
Author: Google Deepmind Agentic AI Team
=============================================================================
"""

import pandas as pd
import numpy as np
import io
import json
from datetime import datetime
from typing import Dict, Any, Tuple, Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

def generate_excel_report(df_clean: pd.DataFrame, profile: Dict[str, Any], audit_log: pd.DataFrame, data_dict: pd.DataFrame) -> bytes:
    """Generates multi-tab OpenPyXL Excel workbook report."""
    wb = openpyxl.Workbook()
    
    # Sheet 1: Executive Summary
    ws_sum = wb.active
    ws_sum.title = "Executive Summary"
    ws_sum.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="1E1E2F", end_color="1E1E2F", fill_type="solid")
    card_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    title_font = Font(name="Segoe UI", size=16, bold=True, color="FFFFFF")

    ws_sum.merge_cells("A1:G2")
    cell_hdr = ws_sum["A1"]
    cell_hdr.value = "THIRANEX SOLUTIONS — DATA QUALITY & CLEANING REPORT"
    cell_hdr.font = title_font
    cell_hdr.fill = header_fill
    cell_hdr.alignment = Alignment(horizontal="center", vertical="center")

    scorecard = profile.get("scorecard", {})
    kpis = [
        ("A4:B5", "OVERALL DATA HEALTH", f"{scorecard.get('overall_score', 0)} / 100", f"Status: {scorecard.get('badge', 'N/A')}"),
        ("C4:D5", "COMPLETENESS SCORE", f"{scorecard.get('completeness_score', 0)}%", f"Missing: {scorecard.get('total_missing', 0)} cells"),
        ("E4:F5", "ACCURACY SCORE", f"{scorecard.get('accuracy_score', 0)}%", f"Outliers: {scorecard.get('total_outliers', 0)}"),
        ("G4:H5", "CONSISTENCY SCORE", f"{scorecard.get('consistency_score', 0)}%", f"Duplicates: {scorecard.get('total_duplicates', 0)}")
    ]

    for rng, title, val, sub in kpis:
        ws_sum.merge_cells(rng)
        top_left = ws_sum[rng.split(":")[0]]
        top_left.value = f"{title}\n\n{val}\n{sub}"
        top_left.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        top_left.fill = card_fill
        top_left.font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")

    # Sheet 2: Cleaned Data
    ws_data = wb.create_sheet(title="Cleaned Data")
    for r in dataframe_to_rows(df_clean, index=False, header=True):
        ws_data.append(r)
    for cell in ws_data[1]:
        cell.fill = header_fill
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")

    # Sheet 3: Audit Log
    if audit_log is not None and not audit_log.empty:
        ws_audit = wb.create_sheet(title="Audit Log")
        for r in dataframe_to_rows(audit_log, index=False, header=True):
            ws_audit.append(r)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

def generate_pdf_report(df_clean: pd.DataFrame, profile: Dict[str, Any], audit_log: pd.DataFrame) -> bytes:
    """Generates ReportLab PDF Executive Briefing."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle("TitleStyle", fontName="Helvetica-Bold", fontSize=18, leading=22, textColor=colors.HexColor("#1F1B2D"))
    body_style = ParagraphStyle("BodyStyle", fontName="Helvetica", fontSize=10, leading=14, textColor=colors.HexColor("#6B6478"))

    elements = [
        Paragraph("<b>THIRANEX SOLUTIONS</b>", ParagraphStyle("Brand", fontName="Helvetica-Bold", fontSize=10, textColor=colors.HexColor("#6D28D9"))),
        Paragraph("Enterprise Data Cleaning & Quality Executive Briefing", title_style),
        Paragraph(f"Generated: {datetime.now().strftime('%B %d, %Y at %H:%M:%S')}", body_style),
        Spacer(1, 10),
        HRFlowable(width="100%", thickness=2, color=colors.HexColor("#6D28D9"), spaceAfter=15)
    ]

    scorecard = profile.get("scorecard", {})
    score_data = [
        ["Overall Score", "Completeness", "Accuracy", "Consistency"],
        [f"{scorecard.get('overall_score', 0)} / 100", f"{scorecard.get('completeness_score', 0)}%", f"{scorecard.get('accuracy_score', 0)}%", f"{scorecard.get('consistency_score', 0)}%"]
    ]
    t_score = Table(score_data, colWidths=[120, 120, 120, 120])
    t_score.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#6D28D9")),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#E5E0F0"))
    ]))
    elements.append(t_score)

    doc.build(elements)
    return buffer.getvalue()

def generate_html_report(df_clean: pd.DataFrame, profile: Dict[str, Any], audit_log: pd.DataFrame) -> str:
    """Generates standalone HTML report with Royal Purple styling."""
    scorecard = profile.get("scorecard", {})
    table_html = df_clean.head(15).to_html(classes="styled-table", index=False)
    
    return f"""<!DOCTYPE html>
<html>
<head>
    <title>Thiranex Solutions — Report</title>
    <style>
        body {{ font-family: 'Inter', -apple-system, sans-serif; background: #F8F7FC; color: #1F1B2D; padding: 30px; }}
        .card {{ background: #FFFFFF; padding: 24px; border-radius: 14px; margin-bottom: 24px; border: 1px solid #E5E0F0; box-shadow: 0 4px 15px rgba(109, 40, 217, 0.05); }}
        h1 {{ color: #1F1B2D; font-size: 24px; margin-top: 0; }}
        h2 {{ color: #6D28D9; font-size: 18px; }}
        p {{ color: #6B6478; font-size: 15px; }}
        .styled-table {{ width: 100%; border-collapse: collapse; border-radius: 8px; overflow: hidden; }}
        .styled-table th {{ background: #6D28D9; color: white; padding: 10px 14px; text-align: left; font-size: 13px; }}
        .styled-table td {{ padding: 10px 14px; border-bottom: 1px solid #E5E0F0; font-size: 13px; color: #1F1B2D; }}
        .styled-table tr:hover td {{ background: rgba(109, 40, 217, 0.03); }}
    </style>
</head>
<body>
    <div class="card">
        <h1>⚡ Thiranex Solutions Data Quality Report</h1>
        <p>Overall Health Score: <b>{scorecard.get('overall_score', 0)} / 100</b></p>
    </div>
    <div class="card">
        <h2>Cleaned Data Preview</h2>
        {table_html}
    </div>
</body>
</html>"""
